import io,json,sys,platform,importlib.metadata as m
import numpy as np,pandas as pd,scipy.optimize as so,matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook,load_workbook
checks=[]
assert np.allclose(np.array([1.,2.])@np.array([3.,4.]),11.); checks.append('numpy_dot')
s=pd.Series([1.,2.,3.],index=pd.date_range('2026-01-01',periods=3,tz='UTC')); assert s.sum()==6 and str(s.index.tz)=='UTC'; checks.append('pandas_timezone_series')
assert abs(so.brentq(lambda x:x*x-2,1,2)**2-2)<1e-10; checks.append('scipy_brentq')
b=io.BytesIO(); fig,ax=plt.subplots(); ax.plot([0,1],[1,2]); fig.savefig(b,format='png'); plt.close(fig); assert b.getvalue().startswith(b'\x89PNG'); checks.append('matplotlib_agg_png')
w=Workbook(); w.active['A1']=2; w.active['A2']='=A1*2'; b=io.BytesIO(); w.save(b); b.seek(0); assert load_workbook(b,data_only=False).active['A2'].value=='=A1*2'; b.seek(0); assert load_workbook(b,data_only=True).active['A2'].value is None; checks.append('openpyxl_roundtrip_formula_not_calculated')
names=['numpy','pandas','scipy','matplotlib','openpyxl','markdown-it-py']
print(json.dumps({'python':sys.version,'platform':platform.platform(),'checks_passed':checks,'distributions':[{ 'name':n,'version':m.version(n),'requires_python':m.metadata(n).get('Requires-Python'),'requires_dist':m.requires(n)} for n in names]},indent=2))
